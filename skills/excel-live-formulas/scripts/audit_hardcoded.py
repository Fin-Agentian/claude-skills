#!/usr/bin/env python3
"""Find hard-coded numbers in an .xlsx that look like they should be formulas.

Usage:
  audit_hardcoded.py <book.xlsx> [--json out.json]

Detects, per sheet:
  - column total : a number equal to the sum of the numbers directly above it (>= 2 cells)
  - row total    : a number equal to the sum of the numbers to its left in the same row (>= 2 cells)
  - share        : a column of fractions equal to (value in another column) / (that column's total)
  - rank         : a column of integers equal to the rank of another numeric column (ties share a rank)
  - number-in-text: text such as "272 人" / "NT$ 1,200" that hides a number (often a computed value)
Nothing is changed; the output is a proposal for the user to accept, extend or skip.
"""
import json
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

EPS = 1e-6
ID_LIKE = re.compile(r"^(\d+|[0-9A-Fa-f-]{8,})$")
NUM_IN_TEXT = re.compile(r"^\s*[^\d\-+]{0,8}?(-?\d[\d,]*(?:\.\d+)?)\s*[^\d]{0,8}\s*$")


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def col_runs(ws):
    """Yield (col, [(row, value), ...]) for maximal vertical runs of numeric constants."""
    for col in range(1, ws.max_column + 1):
        run = []
        for row in range(1, ws.max_row + 2):
            v = ws.cell(row, col).value if row <= ws.max_row else None
            if is_num(v):
                run.append((row, v))
            else:
                if len(run) >= 3:
                    yield col, run
                run = []


def audit(path):
    wb = load_workbook(path)  # formulas are strings, so only constants are numbers
    found = []
    for ws in wb.worksheets:
        def add(kind, ref, value, formula, note):
            found.append({"sheet": ws.title, "cell": ref, "kind": kind, "value": value,
                          "formula": formula, "note": note})

        runs = list(col_runs(ws))
        # column totals: last cell of a run equals the sum of the cells above it
        totals = {}
        for col, run in runs:
            *body, (trow, tval) = run
            if abs(sum(v for _, v in body) - tval) < EPS and len(body) >= 2:
                L = get_column_letter(col)
                add("column total", f"{L}{trow}", tval, f"=SUM({L}{body[0][0]}:{L}{body[-1][0]})",
                    f"equals the sum of {L}{body[0][0]}:{L}{body[-1][0]}")
                totals[col] = (body, trow, tval)

        # row totals: a number equal to the sum of >= 2 numbers to its left
        for row in ws.iter_rows():
            nums = [(c.column, c.value) for c in row if is_num(c.value)]
            for i in range(2, len(nums)):
                col, val = nums[i]
                left = nums[:i]
                for start in range(len(left) - 1):
                    seg = left[start:]
                    if seg[-1][0] != col - 1 or val == 0:
                        continue
                    if any(seg[k + 1][0] != seg[k][0] + 1 for k in range(len(seg) - 1)):
                        continue
                    if abs(sum(v for _, v in seg) - val) < EPS:
                        a, b = get_column_letter(seg[0][0]), get_column_letter(seg[-1][0])
                        add("row total", f"{get_column_letter(col)}{row[0].row}", val,
                            f"=SUM({a}{row[0].row}:{b}{row[0].row})", f"equals {a}:{b} of the same row")
                        break

        # shares and ranks: compare each run with every column that has a total
        for col, run in runs:
            rows = [r for r, _ in run]
            vals = [v for _, v in run]
            L = get_column_letter(col)
            for tcol, (body, trow, tval) in totals.items():
                if tcol == col or tval == 0:
                    continue
                src = dict(body)
                if not all(r in src for r in rows[:len(body)]):
                    continue
                T = get_column_letter(tcol)
                pairs = [(r, v) for r, v in run if r in src]
                if len(pairs) >= 2 and all(abs(v - src[r] / tval) < 1e-4 for r, v in pairs):
                    for r, v in pairs:
                        add("share", f"{L}{r}", v, f"={T}{r}/${T}${trow}", f"{T}{r} as a share of {T}{trow}")
                    break
            if all(isinstance(v, int) or float(v).is_integer() for v in vals) and min(vals) >= 1:
                for tcol, (body, trow, tval) in totals.items():
                    if tcol == col:
                        continue
                    src = dict(body)
                    pairs = [(r, v) for r, v in run if r in src]
                    if len(pairs) < 3:
                        continue
                    ranks = {r: 1 + sum(1 for x in src.values() if x > src[r]) for r, _ in pairs}
                    if all(ranks[r] == v for r, v in pairs):
                        T = get_column_letter(tcol)
                        first, last = body[0][0], body[-1][0]
                        for r, v in pairs:
                            add("rank", f"{L}{r}", v, f"=RANK({T}{r},${T}${first}:${T}${last})",
                                f"rank of {T}{r} within {T}{first}:{T}{last}")
                        break

        # numbers hidden in text ("272 人", "NT$ 1,200")
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("=") and not ID_LIKE.match(c.value.strip()):
                    # 純數字或十六進位字串多半是代碼、ID、雜湊（"04"、"f8251371"），刻意存成文字，不列入
                    m = NUM_IN_TEXT.match(c.value)
                    s, (a, b) = c.value, (m.span(1) if m else (0, 0))
                    glued = (a > 0 and s[a - 1].isascii() and s[a - 1].isalnum()) or \
                            (b < len(s) and s[b].isascii() and s[b].isalnum())  # "C026", "1F" are codes
                    if m and not glued and len(s.replace(m.group(1), "").strip()) <= 6:
                        add("number-in-text", c.coordinate, c.value, None,
                            "store the number (or a formula) and show the unit with a number format")
    # one entry per cell, first match wins
    seen, unique = set(), []
    for f in found:
        key = (f["sheet"], f["cell"])
        if key not in seen:
            seen.add(key)
            unique.append(f)
    return unique


def main(argv):
    if not argv:
        print(__doc__)
        return 2
    items = audit(argv[0])
    if "--json" in argv:
        out = argv[argv.index("--json") + 1]
        json.dump(items, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    by_kind = {}
    for f in items:
        by_kind.setdefault((f["sheet"], f["kind"]), []).append(f)
    print(f"candidates: {len(items)}")
    for (sheet, kind), fs in by_kind.items():
        cells = ", ".join(f["cell"] for f in fs[:12]) + (" ..." if len(fs) > 12 else "")
        print(f"- [{sheet}] {kind} ({len(fs)}): {cells}")
        print(f"    e.g. {fs[0]['cell']} = {fs[0]['value']!r} -> {fs[0]['formula'] or fs[0]['note']}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
