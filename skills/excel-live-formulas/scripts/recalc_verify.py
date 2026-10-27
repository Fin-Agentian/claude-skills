#!/usr/bin/env python3
"""Recalculate an .xlsx with headless LibreOffice and verify its formulas.

Usage:
  recalc_verify.py <book.xlsx> [expected.json] [--keep <recalculated.xlsx>]

- Recalculates a temporary copy (the original file is never modified).
- Reports every formula cell that evaluates to an error (#NAME?, #REF!, #VALUE!, #DIV/0!, ...).
- With expected.json ({"sheet": "<name>", "cells": {"A1": 1, ...}} or a list of those),
  compares each listed cell with the recalculated value (numbers within 1e-6).
- Exit code 0 = all good, 1 = errors or mismatches, 2 = usage / LibreOffice problem.
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ERRORS = ("#NAME?", "#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:", "#SPILL!", "#CALC!")


def find_soffice():
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def recalc(src: Path, workdir: Path) -> Path:
    soffice = find_soffice()
    if not soffice:
        print("LibreOffice (soffice) not found: install LibreOffice Calc, or open the file in Excel to recalculate.")
        sys.exit(2)
    inp = workdir / "in" / src.name
    inp.parent.mkdir()
    shutil.copy(src, inp)
    out_dir = workdir / "out"
    profile = (workdir / "profile").as_uri()
    cmd = [soffice, f"-env:UserInstallation={profile}", "--headless", "--calc",
           "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out_dir), str(inp)]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    out = out_dir / src.name
    if not out.exists():
        print("LibreOffice conversion failed:", res.stderr.strip() or res.stdout.strip())
        sys.exit(2)
    return out


def main(argv):
    args = [a for a in argv if a != "--keep"]
    keep = None
    if "--keep" in argv:
        i = argv.index("--keep")
        keep = Path(argv[i + 1])
        args.remove(argv[i + 1])
    if not args:
        print(__doc__)
        return 2
    src = Path(args[0])
    expected = json.loads(Path(args[1]).read_text(encoding="utf-8")) if len(args) > 1 else []
    if isinstance(expected, dict):
        expected = [expected]

    with tempfile.TemporaryDirectory() as tmp:
        out = recalc(src, Path(tmp))
        if keep:
            shutil.copy(out, keep)
        formulas = load_workbook(src)  # formulas as written
        values = load_workbook(out, data_only=True)  # values after recalculation

        n_formula, problems = 0, []
        for ws in formulas.worksheets:
            vs = values[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        n_formula += 1
                        v = vs[c.coordinate].value
                        # None is a legitimate "" result (e.g. IFERROR(...,"")); only real error codes count
                        if isinstance(v, str) and v.startswith(ERRORS):
                            problems.append(f"{ws.title}!{c.coordinate}: {c.value} -> {v!r}")

        mismatches, checked = [], 0
        for block in expected:
            vs = values[block["sheet"]]
            for ref, want in block["cells"].items():
                checked += 1
                got = vs[ref].value
                if got is None:
                    got = ""
                ok =(isinstance(want, (int, float)) and isinstance(got, (int, float)) and abs(got - want) < 1e-6) \
                    or got == want
                if not ok:
                    mismatches.append(f"{block['sheet']}!{ref}: expected {want!r}, got {got!r}")

    print(f"formula cells: {n_formula}; errors: {len(problems)}; expected checked: {checked}; "
          f"mismatches: {len(mismatches)}")
    for line in (problems + mismatches)[:50]:
        print("  " + line)
    return 1 if problems or mismatches else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
