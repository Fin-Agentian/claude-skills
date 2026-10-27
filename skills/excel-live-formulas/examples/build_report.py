#!/usr/bin/env python3
"""Mode A example: a two-day sales report where every derived number is a formula.

Usage: python3 build_report.py <out.xlsx> <expected.json>
Then:  python3 ../scripts/recalc_verify.py <out.xlsx> <expected.json> --keep <deliver.xlsx>
"""
import json
import random
import sys
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

out_path, expected_path = sys.argv[1:3]
REGIONS = {"01": "North", "02": "South", "03": "East", "04": "West", "05": "Central"}
DAYS = [date(2026, 3, 7), date(2026, 3, 8)]
T = "Orders"

# ---- inputs: synthetic raw records ----
rng = random.Random(7)
orders = [{"id": i, "date": rng.choice(DAYS), "region": rng.choice(list(REGIONS)),
           "customer": f"C{rng.randint(1, 60):03d}", "amount": rng.randint(5, 80) * 10} for i in range(1, 201)]

# ---- expected values, computed in Python only for verification ----
count = {d: Counter(o["region"] for o in orders if o["date"] == d) for d in DAYS}
amount = Counter()
for o in orders:
    amount[o["region"]] += o["amount"]
codes = sorted(REGIONS, key=lambda c: (-amount[c], c))

head = Font(bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="1F6F8B")
param_fill = PatternFill("solid", fgColor="FFF2CC")

wb = Workbook()
ws = wb.active
ws.title = "Summary"
raw = wb.create_sheet("Data")

raw.append(["Id", "Date", "Region", "Customer", "Amount"])
for o in orders:
    raw.append([o["id"], o["date"], o["region"], o["customer"], o["amount"]])
for (d, r) in raw.iter_rows(min_row=2, min_col=2, max_col=3):
    d.number_format, r.number_format = "yyyy-mm-dd", "@"  # real dates; region codes stay text
raw.add_table(Table(displayName=T, ref=f"A1:E{raw.max_row}",
                    tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)))

ws["A1"] = "Sales report"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "Yellow cells are parameters; every other number is a formula over the Data table."
ws.append([])
ws.append(["Rank", "Code", "Region", DAYS[0], DAYS[1], "Orders", "Amount", "Share"])
HR = ws.max_row
for c in ws[HR]:
    c.font, c.fill, c.alignment = head, head_fill, Alignment(horizontal="center")
for col in "DE":  # parameter cells: the header date is the COUNTIFS criterion
    ws[f"{col}{HR}"].number_format = 'm/d "orders"'
    ws[f"{col}{HR}"].fill, ws[f"{col}{HR}"].font = param_fill, Font(bold=True)

FIRST, LAST = HR + 1, HR + len(codes)
TOT = LAST + 1
for i, code in enumerate(codes):
    r = FIRST + i
    ws.append([f"=RANK(G{r},$G${FIRST}:$G${LAST})", code, REGIONS[code],
               f"=COUNTIFS({T}[Date],D${HR},{T}[Region],$B{r})",
               f"=COUNTIFS({T}[Date],E${HR},{T}[Region],$B{r})",
               f"=D{r}+E{r}",
               f"=SUMIFS({T}[Amount],{T}[Region],$B{r})",
               f"=IF($G${TOT}=0,0,G{r}/$G${TOT})"])
    ws[f"B{r}"].number_format = "@"
    ws[f"G{r}"].number_format = '"$"#,##0'
    ws[f"H{r}"].number_format = "0.0%"
ws.append(["", "", "Total"] + [f"=SUM({c}{FIRST}:{c}{LAST})" for c in "DEFGH"])
ws[f"G{TOT}"].number_format = '"$"#,##0'
ws[f"H{TOT}"].number_format = "0.0%"
CUST = TOT + 2  # explicit row: ws.append([]) does not move max_row
ws.cell(CUST, 1, "Distinct customers")
ws.cell(CUST, 4, f"=SUMPRODUCT(1/COUNTIF({T}[Customer],{T}[Customer]))")
ws[f"D{CUST}"].number_format = '0 "customers"'

chart = BarChart()
chart.title = "Orders per region"
chart.add_data(Reference(ws, min_col=4, max_col=5, min_row=HR, max_row=LAST), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=3, min_row=FIRST, max_row=LAST))
ws.add_chart(chart, "J4")
wb.save(out_path)

exp = {}
grand = sum(amount.values())
for i, code in enumerate(codes):
    r = FIRST + i
    exp.update({f"A{r}": 1 + sum(1 for c in codes if amount[c] > amount[code]),
                f"D{r}": count[DAYS[0]][code], f"E{r}": count[DAYS[1]][code],
                f"F{r}": count[DAYS[0]][code] + count[DAYS[1]][code],
                f"G{r}": amount[code], f"H{r}": amount[code] / grand})
exp.update({f"D{TOT}": sum(count[DAYS[0]].values()), f"E{TOT}": sum(count[DAYS[1]].values()),
            f"F{TOT}": len(orders), f"G{TOT}": grand, f"H{TOT}": 1.0,
            f"D{CUST}": len({o["customer"] for o in orders})})
json.dump({"sheet": "Summary", "cells": exp}, open(expected_path, "w", encoding="utf-8"), indent=1)
print(f"wrote {out_path}: {len(exp)} formula cells to verify")
